from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from io import BytesIO

from django.db.models import Count, DecimalField, Q, Sum
from django.db.models.functions import Coalesce
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.lots.models import Lot
from apps.producers.models import Producer
from apps.receptions.models import CoffeeReception


@dataclass(frozen=True)
class DateRange:
    start_date: date | None = None
    end_date: date | None = None


class ReportQueryService:
    @staticmethod
    def receptions_by_producer(*, date_range: DateRange):
        date_filter = Q()
        if date_range.start_date:
            date_filter &= Q(receptions__received_at__date__gte=date_range.start_date)
        if date_range.end_date:
            date_filter &= Q(receptions__received_at__date__lte=date_range.end_date)
        decimal_output = DecimalField(max_digits=14, decimal_places=2)
        return Producer.objects.filter(is_active=True).annotate(
            reception_count=Count("receptions", filter=date_filter, distinct=True),
            received_weight_kg=Coalesce(
                Sum("receptions__calculated_net_weight_kg", filter=date_filter),
                0,
                output_field=decimal_output,
            ),
        ).order_by("full_name")

    @staticmethod
    def lots_summary(*, harvest_year: int | None = None):
        queryset = Lot.objects.annotate(
            origin_count=Count("reception_links", distinct=True),
            producer_count=Count(
                "reception_links__reception__producer",
                distinct=True,
            ),
            event_count=Count("traceability_events", distinct=True),
        )
        if harvest_year:
            queryset = queryset.filter(harvest_year=harvest_year)
        return queryset.order_by("-created_at")

    @staticmethod
    def reception_detail(*, date_range: DateRange):
        queryset = CoffeeReception.objects.select_related("producer", "farm", "received_by")
        if date_range.start_date:
            queryset = queryset.filter(received_at__date__gte=date_range.start_date)
        if date_range.end_date:
            queryset = queryset.filter(received_at__date__lte=date_range.end_date)
        return queryset.order_by("-received_at")


class ReportDataService:
    @staticmethod
    def producer_rows(*, date_range: DateRange) -> list[dict]:
        return [
            {
                "producer_code": producer.code,
                "producer_name": producer.full_name,
                "national_id": producer.national_id,
                "reception_count": producer.reception_count,
                "received_weight_kg": str(producer.received_weight_kg),
            }
            for producer in ReportQueryService.receptions_by_producer(date_range=date_range)
        ]

    @staticmethod
    def lot_rows(*, harvest_year: int | None = None) -> list[dict]:
        return [
            {
                "lot_code": lot.code,
                "lot_name": lot.name,
                "harvest_year": lot.harvest_year,
                "status": lot.get_status_display(),
                "total_weight_kg": str(lot.total_weight_kg),
                "origin_count": lot.origin_count,
                "producer_count": lot.producer_count,
                "event_count": lot.event_count,
            }
            for lot in ReportQueryService.lots_summary(harvest_year=harvest_year)
        ]


class ExcelReportRenderer:
    HEADER_FILL = "3B6B46"

    @classmethod
    def render(cls, *, title: str, headers: list[str], rows: list[list]) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Reporte"
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        title_cell = sheet.cell(row=1, column=1, value=title)
        title_cell.font = Font(size=16, bold=True)
        title_cell.alignment = Alignment(horizontal="center")
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=3, column=column, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=cls.HEADER_FILL)
            cell.alignment = Alignment(horizontal="center")
        for row_index, values in enumerate(rows, start=4):
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row_index, column=column, value=value)
        for column_index in range(1, len(headers) + 1):
            max_length = max(
                len(str(sheet.cell(row=row_index, column=column_index).value or ""))
                for row_index in range(1, sheet.max_row + 1)
            )
            sheet.column_dimensions[get_column_letter(column_index)].width = min(max_length + 3, 45)
        sheet.freeze_panes = "A4"
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()


class PDFReportRenderer:
    @staticmethod
    def render(*, title: str, headers: list[str], rows: list[list]) -> bytes:
        output = BytesIO()
        document = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=12 * mm,
            leftMargin=12 * mm,
            topMargin=12 * mm,
            bottomMargin=12 * mm,
        )
        styles = getSampleStyleSheet()
        elements = [Paragraph(title, styles["Title"]), Spacer(1, 7 * mm)]
        data = [headers, *[[str(value) for value in row] for row in rows]]
        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3B6B46")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D8DFD8")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7F4")]),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        elements.append(table)
        document.build(elements)
        return output.getvalue()


@dataclass(frozen=True, slots=True)
class ReportExportPayload:
    title: str
    headers: list[str]
    rows: list[list]


class ReportExportDataService:
    """Build renderer-neutral export payloads."""

    @classmethod
    def build(
        cls,
        *,
        report_type: str,
        date_range: DateRange,
        harvest_year: int | None,
    ) -> ReportExportPayload:
        if report_type == "lots":
            return cls._lot_payload(harvest_year=harvest_year)
        return cls._producer_payload(date_range=date_range)

    @staticmethod
    def _lot_payload(*, harvest_year: int | None) -> ReportExportPayload:
        data = ReportDataService.lot_rows(harvest_year=harvest_year)
        rows = [
            [
                row["lot_code"],
                row["lot_name"],
                row["harvest_year"],
                row["status"],
                row["total_weight_kg"],
                row["origin_count"],
                row["producer_count"],
                row["event_count"],
            ]
            for row in data
        ]
        return ReportExportPayload(
            title="CoffeeTrace - Reporte por lote",
            headers=[
                "Código",
                "Lote",
                "Cosecha",
                "Estado",
                "Peso kg",
                "Orígenes",
                "Productores",
                "Eventos",
            ],
            rows=rows,
        )

    @staticmethod
    def _producer_payload(*, date_range: DateRange) -> ReportExportPayload:
        data = ReportDataService.producer_rows(date_range=date_range)
        rows = [
            [
                row["producer_code"],
                row["producer_name"],
                row["national_id"],
                row["reception_count"],
                row["received_weight_kg"],
            ]
            for row in data
        ]
        return ReportExportPayload(
            title="CoffeeTrace - Recepciones por productor",
            headers=[
                "Código",
                "Productor",
                "Identificación",
                "Recepciones",
                "Peso kg",
            ],
            rows=rows,
        )
